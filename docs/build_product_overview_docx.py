from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "专注星伴产品总览V1.0.docx"
ICON = ROOT / "content" / "brand" / "app-icon" / "focus-planet-app-icon-512-v1.png"
BUNNY = ROOT / "content" / "brand" / "characters" / "blue-bunny-v1.png"
KITTEN = ROOT / "content" / "brand" / "characters" / "yellow-kitten-v1.png"
PLANET = ROOT / "content" / "brand" / "elements" / "focus-star-planet-v1.png"
DUO = ROOT / "static" / "images" / "runtime" / "brand" / "duo-star-v1.png"
TAB_TRAINING = ROOT / "static" / "images" / "runtime" / "tabbar" / "star-training-v1.png"
TAB_COMPANION = ROOT / "static" / "images" / "runtime" / "tabbar" / "star-companion-v1.png"
TAB_HOME = ROOT / "static" / "images" / "runtime" / "tabbar" / "star-home-v1.png"

BLUE = "2459D8"
BLUE_LIGHT = "EAF1FF"
YELLOW = "FFD35A"
YELLOW_LIGHT = "FFF7DD"
MINT = "55C9A5"
MINT_LIGHT = "EAF8F3"
INK = "1F2A44"
MUTED = "68738D"
GRID = "D9D9D9"
PALE = "F6F8FC"
WHITE = "FFFFFF"
FONT = "Arial Unicode MS"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_cell_margins(cell, top=130, start=150, bottom=130, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def font_run(run, size=11, color=INK, bold=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    return run


def style_paragraph(paragraph, before=0, after=6, line=1.35, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def body(doc, text, bold_lead=None, after=7):
    p = doc.add_paragraph()
    style_paragraph(p, after=after, line=1.45)
    if bold_lead and text.startswith(bold_lead):
        font_run(p.add_run(bold_lead), 11, INK, True)
        font_run(p.add_run(text[len(bold_lead):]), 11, INK)
    else:
        font_run(p.add_run(text), 11, INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=3, line=1.3)
    for r in p.runs:
        font_run(r, 10.5, INK)
    if not p.runs:
        font_run(p.add_run(text), 10.5, INK)
    else:
        p.runs[0].text = text
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4, line=1.3)
    if p.runs:
        p.runs[0].text = text
        font_run(p.runs[0], 10.5, INK)
    else:
        font_run(p.add_run(text), 10.5, INK)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    r = p.add_run(text)
    font_run(r, 18 if level == 1 else 13, "000000", True)
    return p


def centered_image(doc, image_path, width, before=2, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=before, after=after, line=1.0)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=10, line=1.2)
    font_run(p.add_run(text), 9, MUTED)


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.2)
        font_run(p.add_run(header), 10, WHITE, True)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.25)
            font_run(p.add_run(str(value)), 9.5, INK, i == 0)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_tab_structure(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    items = [
        (TAB_TRAINING, "星训练", "每日计划\n五维训练\n20 个可玩游戏", BLUE_LIGHT),
        (TAB_COMPANION, "星陪伴", "番茄节奏\n坐姿陪伴\n休息提醒", MINT_LIGHT),
        (TAB_HOME, "星小屋", "成长报告\n训练历史\n家长设置", YELLOW_LIGHT),
    ]
    for cell, (image, title, copy, fill) in zip(table.rows[0].cells, items):
        set_cell_shading(cell, fill)
        set_cell_border(cell, GRID)
        set_cell_margins(cell, 170, 170, 170, 170)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(0.58))
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p2, before=5, after=5, line=1.1)
        font_run(p2.add_run(title), 12, INK, True)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p3, after=0, line=1.35)
        font_run(p3.add_run(copy), 9.5, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_role_pair(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    roles = [
        (BUNNY, "跃跃", "代表训练、探索、行动与安全守护", BLUE_LIGHT),
        (KITTEN, "暖暖", "代表陪伴、鼓励、休息与成长反馈", YELLOW_LIGHT),
    ]
    for cell, (image, title, copy, fill) in zip(table.rows[0].cells, roles):
        set_cell_shading(cell, fill)
        set_cell_border(cell)
        set_cell_margins(cell, 120, 180, 150, 180)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(1.45))
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p2, before=4, after=4)
        font_run(p2.add_run(title), 13, INK, True)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p3, after=0, line=1.35)
        font_run(p3.add_run(copy), 9.5, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)

    title_style = styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title_style.font.size = Pt(30)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style_p_pr = title_style.element.get_or_add_pPr()
    title_style_border = title_style_p_pr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_p_pr.remove(title_style_border)

    for style_name, size in (("Heading 1", 18), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("专注星伴 产品总览  |  V1.0"), 8.5, MUTED)


def build():
    doc = Document()
    configure_document(doc)

    # Cover
    centered_image(doc, ICON, 2.75, before=8, after=16)
    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, after=8, line=1.1)
    font_run(title.add_run("专注星伴 产品总览"), 30, "000000", True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(sub, after=16, line=1.3)
    font_run(sub.add_run("儿童专注力训练与健康学习陪伴应用"), 14, INK, True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, after=18, line=1.3)
    font_run(meta.add_run("文档版本 V1.0  |  2026年9月"), 10, MUTED)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Inches(0.72)
    intro.paragraph_format.right_indent = Inches(0.72)
    style_paragraph(intro, after=0, line=1.55)
    font_run(intro.add_run("产品已完成跨平台 MVP：20 个专注力训练可玩，星陪伴与端侧观察进入真机验证。本说明用于统一产品定位、能力范围、数据方案、品牌表达与上线前重点。"), 11, INK)

    doc.add_page_break()

    heading(doc, "一 产品定义", 1)
    body(doc, "专注星伴面向 3 至 16 岁儿童与学生，通过短时游戏化训练、番茄节奏、休息提醒、距离与姿态辅助提醒，以及家长可理解的成长报告，帮助孩子逐步形成稳定、舒服、可持续的学习习惯。")
    body(doc, "产品不进行智力、疾病或医学诊断，也不把摄像头信号直接解释为认真或不认真。训练成绩和观察数据用于发现个人趋势、提供温和反馈，而不是评价孩子。")
    centered_image(doc, DUO, 3.0, before=7, after=5)
    add_caption(doc, "跃跃和暖暖陪伴孩子完成训练 专注与休息闭环")

    heading(doc, "二 目标用户与核心场景", 1)
    heading(doc, "儿童与学生", 2)
    make_table(doc, ["年龄", "体验重点"], [
        ("3 至 6 岁", "短时、低文字、强视觉和语音引导"),
        ("7 至 9 岁", "通过明确目标和即时反馈建立训练习惯"),
        ("10 至 12 岁", "提升持续、选择和任务切换能力"),
        ("13 至 16 岁", "用于作业、阅读和阶段复习的自主专注管理"),
    ], [1.25, 5.55])
    heading(doc, "家长", 2)
    bullet(doc, "了解训练是否持续，以及哪些能力维度正在改善。")
    bullet(doc, "查看陪伴时长、休息执行、距离和姿态提醒摘要。")
    bullet(doc, "不需要远程观看孩子，也不需要浏览本地录像。")
    heading(doc, "高频使用场景", 2)
    body(doc, "每日完成约 5 项服务端下发的短训练；写作业、阅读或网课时开启 10 至 25 分钟星陪伴；一轮结束后休息、望远和喝水；家长按日、近 7 天和近 30 天查看变化趋势。")

    heading(doc, "三 核心价值", 1)
    make_table(doc, ["价值", "产品实现"], [
        ("练能力", "用短游戏覆盖五类注意能力，而非只提供倒计时"),
        ("陪过程", "把专注、坐姿自检和休息节奏放进同一个闭环"),
        ("看成长", "强调与孩子过去阶段相比的变化，不做公开排名"),
        ("守隐私", "端侧处理临时相机帧，不录像、不上传原始图片、不做人脸身份识别"),
        ("跨设备连续", "每日计划、完成进度和最佳成绩可由服务端同步"),
    ], [1.35, 5.45])

    doc.add_page_break()
    heading(doc, "四 一级产品结构", 1)
    add_tab_structure(doc)
    body(doc, "星训练负责能力练习，星陪伴负责真实学习过程中的节奏与姿态提醒，星小屋沉淀成长记录和家长管理。三个入口共同形成训练、陪伴、回顾的完整产品闭环。")

    heading(doc, "星训练", 2)
    bullet(doc, "今日训练计划与中断续做。")
    bullet(doc, "五维训练目录、单局过程、暂停退出与结果反馈。")
    bullet(doc, "当前已有 20 个可玩游戏。")
    heading(doc, "星陪伴", 2)
    bullet(doc, "屏幕学习、桌面作业、阅读时间三种模式。")
    bullet(doc, "1 分钟真机快测及 10、15、20、25 分钟常规时长。")
    bullet(doc, "暂停、提前完成、自动休息、喝水和坐姿自检提醒。")
    heading(doc, "星小屋", 2)
    bullet(doc, "成长报告、训练历史、最近陪伴记录。")
    bullet(doc, "儿童档案、家长设置，以及后续的装饰成长和隐私控制。")

    doc.add_page_break()
    heading(doc, "五 五维训练体系", 1)
    centered_image(doc, PLANET, 1.35, before=0, after=5)
    make_table(doc, ["维度", "训练目标", "代表游戏"], [
        ("集中注意", "快速聚焦目标和定位线索", "星星捕手、影子配对、光点追踪、声音在哪里"),
        ("持续注意", "在一段时间内维持稳定投入", "星空守望员、找出变化、小车巡逻、呼吸小星球"),
        ("选择注意", "从干扰中筛选目标信息", "颜色指令、森林寻宝、图案侦探、声音过滤器"),
        ("转换注意", "在规则或任务之间灵活切换", "魔法规则、红绿指令、分类车站、奇偶轨道"),
        ("分配注意", "协调两类以上线索或任务", "双轨任务、记忆与搜索、一心二用小司机、听声找图"),
    ], [1.05, 2.15, 3.6])
    body(doc, "评分应保存正确率、遗漏、误触、反应时间及波动等原始指标，再基于年龄段和个人基线生成成长反馈。MVP 分数仅用于流程验证。")

    doc.add_page_break()
    heading(doc, "六 端侧观察与坐姿辅助", 1)
    body(doc, "前置相机能力默认关闭，由家长主动开启。页面最多每 3 秒分析一帧，原始帧分析后立即释放。", bold_lead="前置相机能力默认关闭")
    make_table(doc, ["阶段", "能力范围"], [
        ("当前可用", "面部是否进入画面；距离是否过近或过远；是否偏离画面中央；头部持续转动、侧倾或俯仰；连续丢失后判断是否离开观察范围"),
        ("下一阶段", "Android 与 iOS 统一引入人体关键点模型，利用耳、肩、髋关键点判断头部侧倾、双肩高低、身体倾斜和明显前倾"),
    ], [1.25, 5.55])
    body(doc, "系统先采集个人自然坐姿基线；异常需连续存在数秒才提醒。关键点不完整时返回无法判断，不得误报为错误坐姿。家长侧只查看稳定在屏比例、过近秒数、姿态提醒次数和休息完成情况等汇总数据。")

    doc.add_page_break()

    heading(doc, "七 数据与服务端能力", 1)
    body(doc, "正式版需要以服务端每日计划为主线，在多设备之间保持训练进度、最佳成绩、报告和资源包状态一致。")
    numbered(doc, "登录后下发每日训练计划并同步多端状态。")
    numbered(doc, "未完成计划跨天保留，完成后可主动切换到新一天。")
    numbered(doc, "支持单题结果重复提交，只保留最佳成绩。")
    numbered(doc, "提供单日或多日报告，以及同一训练上一次成绩对比。")
    numbered(doc, "提供近 7 天、近 30 天、阶段变化和五维能力报告。")
    numbered(doc, "支持资源包清单、后台异步补全、失败重试和低资源量降级循环。")
    heading(doc, "数据原则", 2)
    bullet(doc, "训练结果保留原始指标，报告再做年龄段和个人趋势解释。")
    bullet(doc, "视频不远程查看，原始相机帧不上传。")
    bullet(doc, "儿童档案、摄像头授权、数据删除和导出均由家长控制。")

    heading(doc, "八 品牌与视觉语言", 1)
    add_role_pair(doc)
    centered_image(doc, PLANET, 1.4, before=5, after=3)
    add_caption(doc, "发光星星代表可积累的专注能力 星球轨道代表规律与长期成长")
    body(doc, "跃跃的角色全名为星跃，暖暖的角色全名为星暖，仅在完整角色介绍中使用。两位角色组合称为星伴双星；普通界面和对话优先使用跃跃和暖暖。")
    make_table(doc, ["视觉要素", "表达含义"], [
        ("星际蓝", "训练、探索、稳定与安全"),
        ("阳光黄", "鼓励、成就、温暖与活力"),
        ("薄荷绿", "休息、舒适、健康与陪伴"),
        ("奶油白", "低刺激背景与清楚的信息层级"),
    ], [1.35, 5.45])
    body(doc, "整体表达保持温和、清楚、非评判，避免红色警报、监控感和医学化措辞。")

    doc.add_page_break()

    heading(doc, "九 当前完成度", 1)
    centered_image(doc, DUO, 2.25, before=0, after=6)
    bullet(doc, "3 个一级 Tab 与品牌角色入口已完成。")
    bullet(doc, "20 个规划训练全部可玩，34 个页面已进入 Android 原生编译链路。")
    bullet(doc, "每日训练基础闭环、结果和本地历史可用。")
    bullet(doc, "星陪伴计时、暂停、休息、自检提醒和本地记录可用。")
    bullet(doc, "Android 与 iOS 面部端侧插件已建立，正在进行真机兼容与阈值校准。")
    bullet(doc, "服务端 V1 接口讨论稿已完成。")
    bullet(doc, "App Icon V1 与基础角色素材包已整理。")

    heading(doc, "十 上线前重点", 1)
    numbered(doc, "统一 20 个游戏的图形化教程、暂停、退出与异常恢复。")
    numbered(doc, "完成 Android、iOS 与 HarmonyOS NEXT 真机矩阵和耗电测试。")
    numbered(doc, "校准评分、无效局识别和不同年龄段参数。")
    numbered(doc, "完成家长授权、儿童隐私说明、数据删除导出与第三方 SDK 清单。")
    numbered(doc, "接入服务端每日计划、最佳成绩、报告和资源包补全。")
    numbered(doc, "用真实家庭测试验证提醒频率、误报率、持续使用率和报告可理解性。")

    doc.add_page_break()
    heading(doc, "十一 建议核心指标", 1)
    make_table(doc, ["环节", "核心指标"], [
        ("首次体验", "首次建档完成率；今日计划开始率"),
        ("训练闭环", "今日计划完成率；同一训练阶段稳定性变化"),
        ("持续使用", "次日、7 日和 30 日留存"),
        ("星陪伴", "完整轮次比例；休息完成率"),
        ("端侧观察", "有效帧比例；误报反馈率；关闭率"),
        ("家长价值", "报告查看率；建议理解度"),
    ], [1.35, 5.45])

    doc.add_page_break()

    heading(doc, "十二 产品边界", 1)
    centered_image(doc, ICON, 1.65, before=1, after=8)
    bullet(doc, "不宣称治疗 ADHD 或其他疾病。")
    bullet(doc, "不用单次分数评价智力、性格或学习态度。")
    bullet(doc, "不保存或远程查看儿童视频。")
    bullet(doc, "不使用人脸身份识别、陌生人社交、公开排名或个性化广告。")
    bullet(doc, "涉及摄像头和儿童数据的正式上线版本，必须完成适用地区的法律与儿童隐私合规评审。")
    body(doc, "专注星伴的产品判断标准是：帮助孩子建立可持续的学习习惯，同时让家长获得足够、克制且可理解的信息。")

    doc.core_properties.title = "专注星伴 产品总览"
    doc.core_properties.subject = "儿童专注力训练与健康学习陪伴应用产品说明"
    doc.core_properties.keywords = "专注力训练,儿童学习,星陪伴,坐姿辅助,成长报告"
    doc.core_properties.author = "专注星伴项目组"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
